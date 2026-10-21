import csv
import json

import pytest
from openpyxl import load_workbook

from bndocvqa.services.export import ExportBlocked, export_dataset
from bndocvqa.services.files import sha256_file
from bndocvqa.services.library import LibraryService
from bndocvqa.services.pipeline import Pipeline
from conftest import ITEMS, FakeProvider, make_pdf, run


@pytest.fixture
def two_docs(ctx, doc, tmp_path):
    LibraryService(ctx).import_file(make_pdf(tmp_path / "second.pdf", 1, "B"))  # doc_2, left unprocessed
    Pipeline(ctx).process_document("doc_1")
    run(ctx, FakeProvider([{"questions": ITEMS}, {"questions": ITEMS[:1]}]))
    return ctx


def test_only_finished_documents_exported(two_docs, tmp_path):
    res = export_dataset(two_docs, tmp_path / "out")
    assert res.documents == 1 and res.questions == 4 and res.skipped_unfinished == ["doc_2"]
    assert res.path.name.startswith("bndocvqa_export_")


def test_nothing_finished_blocks(ctx, doc, tmp_path):
    with pytest.raises(ExportBlocked):
        export_dataset(ctx, tmp_path / "out")


def test_files_and_unicode(two_docs, tmp_path):
    root = export_dataset(two_docs, tmp_path / "out").path
    for name in ("index.html", "questions.jsonl", "questions.csv", "questions.xlsx", "manifest.json",
                 "images/doc_1_p1.png", "images/doc_1_p2.png"):
        assert (root / name).is_file(), name
    raw = (root / "questions.jsonl").read_text(encoding="utf-8")
    assert "নোটিশটি কে জারি করেছে?" in raw
    r = json.loads(raw.splitlines()[0])
    assert r["answer_bn"] == r["options"][r["correct_option"]] and r["axis_name"] == "Direct evidence reading"
    assert (root / r["image_path"]).is_file()
    assert (root / "questions.csv").read_bytes()[:3] == b"\xef\xbb\xbf"
    with open(root / "questions.csv", encoding="utf-8-sig", newline="") as f:
        row = next(csv.DictReader(f))
    assert row[f"option_{row['correct_option']}"] == row["answer_bn"]
    ws = load_workbook(root / "questions.xlsx")["questions"]
    assert ws.max_row == 5


def test_html_viewer_embeds_data(two_docs, tmp_path):
    html = (export_dataset(two_docs, tmp_path / "out").path / "index.html").read_text(encoding="utf-8")
    assert "নোটিশটি কে জারি করেছে?" in html and "images/doc_1_p1.png" in html and "<script" in html


def test_manifest_hashes(two_docs, tmp_path):
    root = export_dataset(two_docs, tmp_path / "out").path
    m = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    assert m["question_count"] == 4 and m["documents"][0]["document_id"] == "doc_1"
    for f in m["files"]:
        assert sha256_file(root / f["path"]) == f["sha256"]


def test_ticked_selection_respected(two_docs, tmp_path):
    with pytest.raises(ExportBlocked):
        export_dataset(two_docs, tmp_path / "out", ["doc_2"])


def test_evaluation_on_export(two_docs, tmp_path):
    from bndocvqa.services.evaluation import evaluate, load_jsonl, load_references
    root = export_dataset(two_docs, tmp_path / "out").path
    refs = load_references(root / "questions.jsonl")
    preds = [{"question_id": q, "prediction": r["correct_option"]} for q, r in refs.items()]
    res = evaluate(refs, preds)
    assert res["formats"]["mcq5"]["overall"]["accuracy"] == 1.0
    assert res["formats"]["verify3"]["overall"]["accuracy"] == 1.0


def test_empty_pages_explained_in_export(ctx, doc, tmp_path):
    Pipeline(ctx).process_document("doc_1")
    run(ctx, FakeProvider([{"page_type": "text", "questions": ITEMS[:1]},
                           {"page_type": "signature_or_stamp", "page_note": "seal and signature", "questions": []}]))
    root = export_dataset(ctx, tmp_path / "out").path
    m = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    p2 = m["documents"][0]["pages"][1]
    assert (p2["content_type"], p2["content_note"]) == ("signature_or_stamp", "seal and signature")
    html = (root / "index.html").read_text(encoding="utf-8")
    assert "seal and signature" in html and "No questions" in html


def test_original_documents_included(two_docs, tmp_path):
    import os
    import stat
    root = export_dataset(two_docs, tmp_path / "out").path
    doc = two_docs.documents.get("doc_1")
    copy = root / "documents" / "doc_1.pdf"
    assert copy.is_file() and sha256_file(copy) == doc.file_hash
    assert os.stat(copy).st_mode & stat.S_IWRITE  # the user can move/delete the export
    assert not (root / "documents" / "doc_2.pdf").exists()  # unfinished docs are not exported
    r = json.loads((root / "questions.jsonl").read_text(encoding="utf-8").splitlines()[0])
    assert r["document_path"] == "documents/doc_1.pdf" and (root / r["document_path"]).is_file()
    m = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    assert m["documents"][0]["document_path"] == "documents/doc_1.pdf"
    assert any(f["path"] == "documents/doc_1.pdf" for f in m["files"])
    assert '"doc_1": "documents/doc_1.pdf"' in (root / "index.html").read_text(encoding="utf-8")
